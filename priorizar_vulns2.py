#!/usr/bin/env python3
"""
Priorização de vulnerabilidades do Wazuh (Vulnerability Detection).

Enriquece o CSV exportado do Wazuh com:
  - EPSS (FIRST)            -> probabilidade de exploração em 30 dias
  - CISA KEV                -> exploração confirmada no mundo real
  - Contexto do ativo       -> criticidade / exposição / ambiente (CSV opcional)
  - Histórico               -> variação do EPSS, CVEs novos/resolvidos, tendência

Saídas (em --saida-dir):
  priorizacao_AAAA-MM-DD.xlsx   abas: Resumo, Por CVE, Remediação, Por Ativo, Detalhe, Tendência
  priorizacao_AAAA-MM-DD.csv    detalhe (uma linha por agente x pacote x CVE) para automação
  priorizacao_AAAA-MM-DD.json   visão por CVE para integrações
  relatorio_vulnerabilidades_AAAA-MM-DD.html  relatório executivo/técnico
  historico/cves_AAAA-MM-DD.csv snapshot usado para calcular tendência e variação do EPSS

Uso:
  python priorizar_vulns.py
  python priorizar_vulns.py --entrada input/vulns2.csv --ativos input/ativos.csv
"""

import argparse
import csv
import gzip
import json
import re
import time
import unicodedata
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
import requests
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList

# =============================================================================
# CONFIGURAÇÕES
# =============================================================================

# Fontes de dados. Confira as URLs no site do EPSS/CISA se algum download falhar.
EPSS_BULK_URLS = [
    "https://epss.empiricalsecurity.com/epss_scores-current.csv.gz",
    "https://epss.cyentia.com/epss_scores-current.csv.gz",
]
EPSS_API = "https://api.first.org/data/v1/epss"
KEV_URL = (
    "https://www.cisa.gov/sites/default/files/feeds/"
    "known_exploited_vulnerabilities.json"
)

# Limites da API EPSS (usada só como fallback se o download em massa falhar)
LIMITE_CARACTERES = 1800
MAX_TENTATIVAS = 3

# Regras das faixas. CALIBRE ao seu ambiente.
#   P0: está no CISA KEV
#   P1: EPSS >= p1_epss  OU  percentil >= p1_percentil
#       OU (CVSS >= p1_cvss_alto E EPSS >= p1_cvss_alto_epss)
#   P2: CVSS >= p2_cvss  OU  percentil >= p2_percentil
#   P3: demais
REGRAS = {
    "p1_epss": 0.50,
    "p1_percentil": 0.95,
    "p1_cvss_alto": 9.0,
    "p1_cvss_alto_epss": 0.10,
    "p2_cvss": 7.0,
    "p2_percentil": 0.80,
    "epss_salto": 0.10,  # variação de EPSS que vira alerta no "motivo"
}

# Contexto do ativo: cada valor soma/subtrai níveis (positivo = mais urgente).
# O ajuste total é limitado a +1 / -1 nível. Só o KEV pode ser P0: um ativo
# crítico nunca promove um CVE não-KEV para P0, e um CVE KEV nunca cai abaixo de P1.
AJUSTE_EXPOSICAO = {"internet": 1, "dmz": 0, "interna": 0, "isolada": -1}
AJUSTE_CRITICIDADE = {"critica": 1, "alta": 0, "media": 0, "baixa": -1}
AJUSTE_AMBIENTE = {
    "producao": 0,
    "homologacao": -1,
    "teste": -1,
    "desenvolvimento": -1,
    "dev": -1,
}

NIVEIS = {0: "P0", 1: "P1", 2: "P2", 3: "P3"}
DESCRICAO_FAIXA = {
    "P0": "Exploração confirmada (CISA KEV). Corrigir/mitigar imediatamente.",
    "P1": "Exploração provável (EPSS alto) ou crítico com EPSS relevante.",
    "P2": "CVSS alto ou EPSS acima do percentil 80. Janela de patch normal.",
    "P3": "Demais. Tratar no ciclo regular de atualização.",
}

# Nomes das colunas do CSV do Wazuh (aceita mais de um nome por campo)
ALIASES = {
    "cve": ["vulnerability.id"],
    "cvss": ["vulnerability.score.base"],
    "agente": ["agent.name"],
    "pacote": ["package.name"],
    "versao": ["package.version"],
    "severidade_wazuh": ["vulnerability.severity"],
    "publicado": ["vulnerability.published_at", "vulnerability.published"],
    "descricao": ["vulnerability.description"],
    "condicao": ["vulnerability.scanner.condition"],
}
OBRIGATORIOS = ["cve", "cvss"]

# =============================================================================
# UTILITÁRIOS
# =============================================================================


def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}")


def erro_curto(erro, limite=150):
    texto = str(erro).replace("\n", " ")
    return texto if len(texto) <= limite else texto[:limite] + "..."


def normalizar(valor):
    """minúsculas, sem acento, sem espaços nas pontas."""
    if valor is None or (isinstance(valor, float) and np.isnan(valor)):
        return ""
    s = unicodedata.normalize("NFKD", str(valor))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().lower()


def detectar_separador(caminho):
    with open(caminho, encoding="utf-8-sig", errors="replace") as f:
        primeira = f.readline()
    candidatos = [",", ";", "\t", "|"]
    return max(candidatos, key=primeira.count)


def baixar(url, destino, timeout=120):
    tmp = destino.with_suffix(destino.suffix + ".part")
    with requests.get(url, stream=True, timeout=timeout) as r:
        r.raise_for_status()
        with open(tmp, "wb") as f:
            for pedaco in r.iter_content(1 << 20):
                f.write(pedaco)
    tmp.replace(destino)


# =============================================================================
# ENTRADA: CSV DO WAZUH
# =============================================================================


def carregar_wazuh(caminho, avisos):
    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    log(f"Lendo {caminho} ...")
    bruto = pd.read_csv(caminho, sep=detectar_separador(caminho), dtype=str)
    log(f"Registros no CSV: {len(bruto)}")

    df = pd.DataFrame(index=bruto.index)
    ausentes_opc = []
    for interno, candidatos in ALIASES.items():
        achada = next((c for c in candidatos if c in bruto.columns), None)
        if achada is None:
            if interno in OBRIGATORIOS:
                raise ValueError(
                    f"Coluna '{candidatos[0]}' não encontrada. "
                    f"Colunas disponíveis: {list(bruto.columns)}"
                )
            ausentes_opc.append(candidatos[0])
            df[interno] = np.nan
        else:
            df[interno] = bruto[achada]

    if ausentes_opc:
        avisos.append(
            "Colunas opcionais ausentes no CSV (campos ficam vazios): "
            + ", ".join(ausentes_opc)
        )

    df["cve"] = df["cve"].str.strip().str.upper()
    vazios = df["cve"].isna() | (df["cve"] == "")
    if vazios.any():
        avisos.append(f"{int(vazios.sum())} linhas sem vulnerability.id foram descartadas.")
        df, bruto = df[~vazios], bruto[~vazios]

    df["cvss"] = pd.to_numeric(df["cvss"], errors="coerce")
    df["agente"] = df["agente"].fillna("(desconhecido)").str.strip()
    df["pacote"] = df["pacote"].fillna("(desconhecido)").str.strip()
    df["versao"] = df["versao"].fillna("(desconhecido)").str.strip()

    df["cve_valido"] = df["cve"].str.fullmatch(r"CVE-\d{4}-\d{4,}", na=False)
    invalidos = df.loc[~df["cve_valido"], "cve"]
    if len(invalidos):
        exemplos = ", ".join(invalidos.unique()[:5])
        avisos.append(
            f"{len(invalidos)} registros com ID fora do padrão CVE (ex.: {exemplos}). "
            "Mantidos na saída, mas sem EPSS/KEV."
        )

    sem_cvss = int(df["cvss"].isna().sum())
    if sem_cvss:
        avisos.append(f"{sem_cvss} registros sem CVSS (priorizados só por KEV/EPSS).")

    # colunas originais do Wazuh seguem no Detalhe, depois das colunas de análise
    return df, bruto


# =============================================================================
# EPSS
# =============================================================================


def ler_epss_bulk(caminho):
    with gzip.open(caminho, "rt", encoding="utf-8") as f:
        primeira = f.readline()
    m = re.search(r"score_date:(\d{4}-\d{2}-\d{2})", primeira)
    data = m.group(1) if m else None

    tabela = pd.read_csv(caminho, compression="gzip", comment="#")
    tabela.columns = [c.strip().lower() for c in tabela.columns]
    tabela["cve"] = tabela["cve"].str.upper()
    tabela = tabela.rename(columns={"percentile": "epss_percentil"})
    tabela = tabela[["cve", "epss", "epss_percentil"]].drop_duplicates("cve")
    return tabela.set_index("cve"), data


def epss_via_api(cves, avisos):
    """Fallback: consulta em lotes à API do FIRST."""
    lotes, atual = [], []
    for cve in cves:
        candidato = atual + [cve]
        if len(",".join(candidato)) > LIMITE_CARACTERES and atual:
            lotes.append(atual)
            atual = [cve]
        else:
            atual = candidato
    if atual:
        lotes.append(atual)

    linhas, falhos, data_ref = [], 0, None
    with requests.Session() as sessao:
        for n, lote in enumerate(lotes, start=1):
            log(f"EPSS (API) lote {n}/{len(lotes)} ...")
            for tentativa in range(1, MAX_TENTATIVAS + 1):
                try:
                    r = sessao.get(EPSS_API, params={"cve": ",".join(lote)}, timeout=30)
                    r.raise_for_status()
                    payload = r.json()
                    if payload.get("status") != "OK":
                        raise ValueError(f"status inesperado: {payload.get('status')}")
                    for item in payload.get("data", []):
                        linhas.append(
                            (item["cve"].upper(), float(item["epss"]), float(item["percentile"]))
                        )
                        data_ref = item.get("date", data_ref)
                    break
                except (requests.RequestException, ValueError) as erro:
                    log(f"  lote {n}, tentativa {tentativa}: {erro_curto(erro)}")
                    if tentativa == MAX_TENTATIVAS:
                        falhos += len(lote)
                    else:
                        time.sleep(2 * tentativa)
            time.sleep(0.2)

    if falhos:
        avisos.append(
            f"EPSS (API): {falhos} CVEs em lotes que falharam. Eles aparecem como 'sem EPSS', "
            "não como baixo risco."
        )
    tabela = pd.DataFrame(linhas, columns=["cve", "epss", "epss_percentil"])
    return tabela.drop_duplicates("cve").set_index("cve"), data_ref


def obter_epss(cves, cache_dir, hoje, avisos):
    """Retorna (tabela indexada por CVE, data do score, origem)."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_hoje = cache_dir / f"epss_{hoje}.csv.gz"

    # cache de hoje corrompido -> descarta
    if cache_hoje.exists():
        try:
            ler_epss_bulk(cache_hoje)
        except Exception:
            cache_hoje.unlink()

    if not cache_hoje.exists():
        for url in EPSS_BULK_URLS:
            try:
                log(f"Baixando EPSS completo: {url}")
                baixar(url, cache_hoje)
                ler_epss_bulk(cache_hoje)  # valida
                break
            except Exception as erro:
                log(f"  falhou: {erro_curto(erro)}")
                cache_hoje.unlink(missing_ok=True)

    arquivo = cache_hoje if cache_hoje.exists() else None
    if arquivo is None:
        anteriores = sorted(cache_dir.glob("epss_*.csv.gz"))
        if anteriores:
            arquivo = anteriores[-1]
            avisos.append(f"EPSS: download falhou; usando cache antigo ({arquivo.name}).")

    if arquivo is not None:
        tabela, data = ler_epss_bulk(arquivo)
        data = data or arquivo.name.split("_")[1].split(".")[0]
        return tabela, data, f"arquivo completo ({arquivo.name})"

    log("Download em massa indisponível. Usando a API do FIRST ...")
    tabela, data = epss_via_api(cves, avisos)
    return tabela, data, "API FIRST"


# =============================================================================
# CISA KEV
# =============================================================================


def obter_kev(cache_dir, hoje, avisos):
    """Retorna (DataFrame indexado por CVE, disponivel: bool)."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_hoje = cache_dir / f"kev_{hoje}.json"

    if not cache_hoje.exists():
        try:
            log("Baixando CISA KEV ...")
            baixar(KEV_URL, cache_hoje, timeout=60)
            json.loads(cache_hoje.read_text(encoding="utf-8"))  # valida
        except Exception as erro:
            log(f"  falhou: {erro_curto(erro)}")
            cache_hoje.unlink(missing_ok=True)

    arquivo = cache_hoje if cache_hoje.exists() else None
    if arquivo is None:
        anteriores = sorted(cache_dir.glob("kev_*.json"))
        if anteriores:
            arquivo = anteriores[-1]
            avisos.append(f"KEV: download falhou; usando cache antigo ({arquivo.name}).")

    if arquivo is None:
        avisos.append(
            "KEV INDISPONÍVEL: nenhum CVE foi marcado como P0. O resultado está incompleto. "
            "Rode novamente com acesso à internet."
        )
        return pd.DataFrame(), False

    itens = json.loads(arquivo.read_text(encoding="utf-8")).get("vulnerabilities", [])
    kev = pd.DataFrame(
        {
            "cve": [i["cveID"].upper() for i in itens],
            "kev_data_inclusao": [i.get("dateAdded") for i in itens],
            "kev_prazo": [i.get("dueDate") for i in itens],
            "kev_ransomware": [
                str(i.get("knownRansomwareCampaignUse", "")).lower() == "known" for i in itens
            ],
            "kev_nome": [i.get("vulnerabilityName") for i in itens],
        }
    ).drop_duplicates("cve")
    return kev.set_index("cve"), True


# =============================================================================
# CONTEXTO DO ATIVO
# =============================================================================


def carregar_ativos(caminho, avisos):
    if caminho is None or not caminho.exists():
        avisos.append(
            f"Contexto de ativos não utilizado ({caminho} não existe). Sem ele, a priorização "
            "ignora criticidade, exposição e ambiente."
        )
        return None

    a = pd.read_csv(caminho, sep=detectar_separador(caminho), dtype=str)
    a.columns = [normalizar(c) for c in a.columns]
    col_agente = next((c for c in ("agent.name", "agente", "agent_name") if c in a.columns), None)
    if col_agente is None:
        raise ValueError(
            f"{caminho}: coluna 'agent.name' não encontrada (colunas: {list(a.columns)})."
        )
    for c in ("criticidade", "exposicao", "ambiente"):
        if c not in a.columns:
            a[c] = ""
    a = a.rename(columns={col_agente: "agente"})
    a["_chave"] = a["agente"].map(normalizar)
    a = a.drop_duplicates("_chave")

    validos = {
        "criticidade": AJUSTE_CRITICIDADE,
        "exposicao": AJUSTE_EXPOSICAO,
        "ambiente": AJUSTE_AMBIENTE,
    }
    for campo, mapa in validos.items():
        desconhecidos = {
            v for v in a[campo].dropna().unique() if normalizar(v) and normalizar(v) not in mapa
        }
        if desconhecidos:
            avisos.append(
                f"ativos.csv: valores não reconhecidos em '{campo}' (ignorados): "
                f"{', '.join(sorted(desconhecidos))}. Aceitos: {', '.join(mapa)}."
            )
    return a[["_chave", "criticidade", "exposicao", "ambiente"]]


def _ajuste_contexto(crit, expo, amb):
    partes, total = [], 0
    for rotulo, valor, mapa in (
        ("criticidade", crit, AJUSTE_CRITICIDADE),
        ("exposição", expo, AJUSTE_EXPOSICAO),
        ("ambiente", amb, AJUSTE_AMBIENTE),
    ):
        pts = mapa.get(normalizar(valor), 0)
        if pts:
            total += pts
            partes.append(f"{rotulo} {valor} ({pts:+d})")
    total = max(-1, min(1, total))
    texto = f"contexto do ativo {total:+d}: " + ", ".join(partes) if total else ""
    return total, texto


# =============================================================================
# HISTÓRICO
# =============================================================================


def carregar_snapshot_anterior(hist_dir, hoje):
    if not hist_dir.exists():
        return None, None
    arquivos = sorted(
        p for p in hist_dir.glob("cves_*.csv") if p.stem.replace("cves_", "") < hoje
    )
    if not arquivos:
        return None, None
    try:
        snap = pd.read_csv(arquivos[-1])
        return snap.set_index("cve"), arquivos[-1].stem.replace("cves_", "")
    except Exception:
        return None, None


def montar_tendencia(hist_dir):
    linhas = []
    for p in sorted(hist_dir.glob("cves_*.csv")):
        try:
            s = pd.read_csv(p)
            cont = s["prioridade"].value_counts()
            linhas.append(
                {
                    "data": p.stem.replace("cves_", ""),
                    "P0": int(cont.get("P0", 0)),
                    "P1": int(cont.get("P1", 0)),
                    "P2": int(cont.get("P2", 0)),
                    "P3": int(cont.get("P3", 0)),
                    "total_cves": len(s),
                }
            )
        except Exception:
            continue
    return pd.DataFrame(linhas, columns=["data", "P0", "P1", "P2", "P3", "total_cves"])


# =============================================================================
# PRIORIZAÇÃO
# =============================================================================


def montar_motivo(kev, ransom, prazo, epss, pct, cvss, delta, ajuste_txt):
    p = []
    if kev:
        s = "KEV (exploração confirmada)"
        if ransom:
            s += " + ransomware"
        if isinstance(prazo, str) and prazo:
            s += f", prazo CISA {prazo}"
        p.append(s)
    if pd.notna(epss):
        if epss >= 0.05 or (pd.notna(pct) and pct >= 0.80):
            p.append(f"EPSS {epss:.2f} (p{pct * 100:.0f})")
    else:
        p.append("sem EPSS")
    if pd.notna(cvss):
        if cvss >= 7:
            p.append(f"CVSS {cvss:.1f}")
    else:
        p.append("sem CVSS")
    if pd.notna(delta) and delta >= REGRAS["epss_salto"]:
        p.append(f"EPSS subiu {delta:+.2f} desde a última execução")
    if ajuste_txt:
        p.append(ajuste_txt)
    return "; ".join(p) if p else "sem sinais de exploração"


def priorizar(df, epss, kev, kev_ok, ativos, anterior, hoje_dt):
    # --- EPSS / KEV ---
    df = df.join(epss, on="cve")
    df.loc[~df["cve_valido"], ["epss", "epss_percentil"]] = np.nan
    df["sem_epss"] = df["epss"].isna()

    if kev_ok:
        df = df.join(kev, on="cve")
    else:
        for c in ("kev_data_inclusao", "kev_prazo", "kev_nome"):
            df[c] = np.nan
        df["kev_ransomware"] = False
    df["in_kev"] = df["kev_data_inclusao"].notna() if kev_ok else False
    df["kev_ransomware"] = df["kev_ransomware"].fillna(False).astype(bool)

    # --- histórico ---
    if anterior is not None:
        df["epss_anterior"] = df["cve"].map(anterior["epss"])
        df["epss_delta"] = df["epss"] - df["epss_anterior"]
    else:
        df["epss_anterior"] = np.nan
        df["epss_delta"] = np.nan

    # --- contexto do ativo ---
    df["_chave"] = df["agente"].map(normalizar)
    if ativos is not None:
        df = df.merge(ativos, on="_chave", how="left")
    else:
        for c in ("criticidade", "exposicao", "ambiente"):
            df[c] = ""
    for c in ("criticidade", "exposicao", "ambiente"):
        df[c] = df[c].fillna("")

    combos = df[["criticidade", "exposicao", "ambiente"]].drop_duplicates().copy()
    resultado = [
        _ajuste_contexto(c, e, a)
        for c, e, a in zip(combos["criticidade"], combos["exposicao"], combos["ambiente"])
    ]
    combos["ajuste_contexto"] = [r[0] for r in resultado]
    combos["_ajuste_txt"] = [r[1] for r in resultado]
    df = df.merge(combos, on=["criticidade", "exposicao", "ambiente"], how="left")

    # --- faixa base (vetorizado) ---
    cvss, epss_v, pct = df["cvss"], df["epss"], df["epss_percentil"]
    p1 = (
        (epss_v >= REGRAS["p1_epss"])
        | (pct >= REGRAS["p1_percentil"])
        | ((cvss >= REGRAS["p1_cvss_alto"]) & (epss_v >= REGRAS["p1_cvss_alto_epss"]))
    )
    p2 = (cvss >= REGRAS["p2_cvss"]) | (pct >= REGRAS["p2_percentil"])
    base = np.select([df["in_kev"], p1, p2], [0, 1, 2], default=3)

    minimo = np.where(df["in_kev"], 0, 1)
    df["nivel"] = np.clip(base - df["ajuste_contexto"], minimo, 3).astype(int)
    df["prioridade"] = df["nivel"].map(NIVEIS)

    # desempate dentro da faixa (NÃO define a faixa)
    df["score_desempate"] = (
        40 * cvss.fillna(0) / 10 + 50 * pct.fillna(0) + 10 * df["kev_ransomware"]
    ).round(2)

    df["motivo_prioridade"] = [
        montar_motivo(k, r, pz, e, pc, cv, d, t)
        for k, r, pz, e, pc, cv, d, t in zip(
            df["in_kev"],
            df["kev_ransomware"],
            df["kev_prazo"],
            df["epss"],
            df["epss_percentil"],
            df["cvss"],
            df["epss_delta"],
            df["_ajuste_txt"],
        )
    ]

    pub = pd.to_datetime(df["publicado"], errors="coerce", utc=True).dt.tz_localize(None)
    df["idade_dias"] = (hoje_dt - pub).dt.days
    df["publicado"] = pub.dt.strftime("%Y-%m-%d")

    return df.sort_values(["nivel", "score_desempate"], ascending=[True, False]).reset_index(
        drop=True
    )


# =============================================================================
# VISÕES
# =============================================================================


def _lista_curta(serie, limite=10):
    itens = sorted(set(serie.dropna()))
    texto = ", ".join(itens[:limite])
    if len(itens) > limite:
        texto += f" (+{len(itens) - limite})"
    return texto


def visao_por_cve(df, anterior):
    melhor = df.drop_duplicates("cve")  # df já vem ordenado: 1ª linha = pior caso
    contagens = df.groupby("cve").agg(
        agentes_afetados=("agente", "nunique"),
        pacotes_afetados=("pacote", "nunique"),
    )
    por_cve = melhor.join(contagens, on="cve").rename(columns={"agente": "ativo_mais_critico"})
    por_cve["descricao"] = por_cve["descricao"].fillna("").str.slice(0, 300)

    if anterior is not None:
        por_cve["novo"] = ~por_cve["cve"].isin(anterior.index)
    else:
        por_cve["novo"] = False

    colunas = [
        "cve", "prioridade", "motivo_prioridade", "cvss", "severidade_wazuh",
        "epss", "epss_percentil", "epss_delta", "sem_epss",
        "in_kev", "kev_ransomware", "kev_prazo", "kev_data_inclusao",
        "agentes_afetados", "pacotes_afetados", "ativo_mais_critico",
        "publicado", "idade_dias", "novo", "score_desempate", "descricao",
    ]
    return por_cve[colunas].reset_index(drop=True)


def visao_remediacao(df):
    chave = ["pacote", "versao"]
    por_cve_pkg = df.groupby(chave + ["cve"], as_index=False)["nivel"].min()

    cont = (
        por_cve_pkg.pivot_table(
            index=chave, columns="nivel", values="cve", aggfunc="nunique", fill_value=0
        )
        .reindex(columns=[0, 1, 2, 3], fill_value=0)
        .rename(columns=NIVEIS)
        .reset_index()
    )
    cont.columns.name = None

    hosts = df.groupby(chave).agg(
        hosts_afetados=("agente", "nunique"),
        lista_hosts=("agente", _lista_curta),
        condicao_correcao=("condicao", "first"),
    ).reset_index()

    criticos = (
        por_cve_pkg[por_cve_pkg["nivel"] <= 1]
        .sort_values(["nivel", "cve"])
        .groupby(chave)["cve"]
        .agg(lambda s: ", ".join(list(s)[:5]) + (f" (+{len(s) - 5})" if len(s) > 5 else ""))
        .rename("cves_p0_p1")
        .reset_index()
    )
    nivel_max = por_cve_pkg.groupby(chave)["nivel"].min().rename("_nivel").reset_index()

    rem = cont.merge(hosts, on=chave).merge(nivel_max, on=chave).merge(criticos, on=chave, how="left")
    rem["cves_eliminados"] = rem[["P0", "P1", "P2", "P3"]].sum(axis=1)
    rem["prioridade_max"] = rem["_nivel"].map(NIVEIS)
    rem["cves_p0_p1"] = rem["cves_p0_p1"].fillna("")
    rem = rem.sort_values(
        ["_nivel", "P0", "P1", "cves_eliminados"], ascending=[True, False, False, False]
    )
    rem = rem.rename(columns={"versao": "versao_instalada"})
    return rem[
        [
            "pacote", "versao_instalada", "prioridade_max", "hosts_afetados", "cves_eliminados",
            "P0", "P1", "P2", "P3", "cves_p0_p1", "lista_hosts", "condicao_correcao",
        ]
    ].reset_index(drop=True)


def visao_por_ativo(df):
    por_cve_ag = df.groupby(["agente", "cve"], as_index=False)["nivel"].min()
    cont = (
        por_cve_ag.pivot_table(
            index="agente", columns="nivel", values="cve", aggfunc="nunique", fill_value=0
        )
        .reindex(columns=[0, 1, 2, 3], fill_value=0)
        .rename(columns=NIVEIS)
    )
    cont.columns.name = None
    extra = df.groupby("agente").agg(
        criticidade=("criticidade", "first"),
        exposicao=("exposicao", "first"),
        ambiente=("ambiente", "first"),
        pacotes_vulneraveis=("pacote", "nunique"),
    )
    ativo = cont.join(extra)
    ativo["cves_total"] = ativo[["P0", "P1", "P2", "P3"]].sum(axis=1)
    ativo = ativo.reset_index().sort_values(
        ["P0", "P1", "P2", "cves_total"], ascending=False
    )
    return ativo[
        [
            "agente", "criticidade", "exposicao", "ambiente",
            "P0", "P1", "P2", "P3", "cves_total", "pacotes_vulneraveis",
        ]
    ].reset_index(drop=True)


def visao_detalhe(df, originais):
    analise = [
        "prioridade", "motivo_prioridade", "score_desempate", "cve", "agente", "pacote",
        "versao", "cvss", "epss", "epss_percentil", "epss_delta", "sem_epss",
        "in_kev", "kev_ransomware", "kev_prazo", "criticidade", "exposicao", "ambiente",
        "ajuste_contexto",
    ]
    det = df[analise].copy()
    # junta as colunas originais do Wazuh (mesmo índice original preservado em _idx)
    orig = originais.loc[df["_idx"]].reset_index(drop=True)
    orig = orig.add_prefix("wazuh | ")
    return pd.concat([det.reset_index(drop=True), orig], axis=1)


# =============================================================================
# EXCEL
# =============================================================================

COR_FAIXA = {
    "P0": ("C00000", "FFFFFF"),
    "P1": ("F4B183", "000000"),
    "P2": ("FFE699", "000000"),
    "P3": ("C6E0B4", "000000"),
}
FORMATOS = {
    "epss": "0.0000",
    "epss_percentil": "0.0%",
    "epss_delta": "+0.0000;-0.0000;0.0000",
    "epss_anterior": "0.0000",
    "cvss": "0.0",
    "score_desempate": "0.00",
}
COLUNAS_BOOL = ("in_kev", "kev_ransomware", "sem_epss", "novo")

FONTE = "Arial"


def _para_excel(df):
    df = df.copy()
    for c in COLUNAS_BOOL:
        if c in df.columns:
            df[c] = df[c].map({True: "Sim", False: ""})
    return df


def _formatar_aba(ws, df, largura_max=60):
    cab_fonte = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
    cab_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = cab_fonte
        cell.fill = cab_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions

    amostra = df.head(500)
    for i, col in enumerate(df.columns, start=1):
        maior = max([len(str(col))] + [len(str(v)) for v in amostra[col].tolist()])
        ws.column_dimensions[get_column_letter(i)].width = min(max(maior + 2, 8), largura_max)
        formato = FORMATOS.get(col)
        if formato:
            for linha in range(2, len(df) + 2):
                ws.cell(row=linha, column=i).number_format = formato

    if "prioridade" in df.columns and len(df):
        letra = get_column_letter(list(df.columns).index("prioridade") + 1)
        faixa = f"{letra}2:{letra}{len(df) + 1}"
    else:
        faixa = None
    if faixa:
        for nome, (fundo, texto) in COR_FAIXA.items():
            ws.conditional_formatting.add(
                faixa,
                CellIsRule(
                    operator="equal",
                    formula=[f'"{nome}"'],
                    fill=PatternFill(start_color=fundo, end_color=fundo, fill_type="solid"),
                    font=Font(name=FONTE, bold=True, color=texto),
                ),
            )
    # colunas P0..P3 (Remediação / Por Ativo): destaca quando > 0
    for nome in ("P0", "P1"):
        if nome in df.columns and len(df):
            letra = get_column_letter(list(df.columns).index(nome) + 1)
            fundo, texto = COR_FAIXA[nome]
            ws.conditional_formatting.add(
                f"{letra}2:{letra}{len(df) + 1}",
                CellIsRule(
                    operator="greaterThan",
                    formula=["0"],
                    fill=PatternFill(start_color=fundo, end_color=fundo, fill_type="solid"),
                    font=Font(name=FONTE, bold=True, color=texto),
                ),
            )


def _montar_resumo(ws, ctx, por_cve, detalhe):
    negrito = Font(name=FONTE, bold=True, size=10)
    normal = Font(name=FONTE, size=10)
    cab_fonte = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
    cab_fill = PatternFill("solid", fgColor="1F3864")
    borda = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16

    ws["A1"] = "Priorização de vulnerabilidades (Wazuh + EPSS + CISA KEV)"
    ws["A1"].font = Font(name=FONTE, bold=True, size=14, color="1F3864")

    info = [
        ("Gerado em", ctx["gerado_em"]),
        ("Arquivo de entrada", str(ctx["entrada"])),
        ("EPSS", ctx["epss_info"]),
        ("CISA KEV", ctx["kev_info"]),
        ("Contexto de ativos", ctx["ativos_info"]),
        ("Execução anterior", ctx["anterior_info"]),
    ]
    linha = 3
    for rotulo, valor in info:
        ws.cell(row=linha, column=1, value=rotulo).font = negrito
        c = ws.cell(row=linha, column=2, value=valor)
        c.font = normal
        c.alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1

    # --- tabela de faixas (com fórmulas ligadas às abas de dados) ---
    linha += 1
    for j, titulo in enumerate(["Faixa", "Critério", "CVEs únicos", "Ocorrências"], start=1):
        c = ws.cell(row=linha, column=j, value=titulo)
        c.font, c.fill, c.border = cab_fonte, cab_fill, borda
    inicio = linha + 1

    col_prio_cve = get_column_letter(list(por_cve.columns).index("prioridade") + 1)
    col_prio_det = get_column_letter(list(detalhe.columns).index("prioridade") + 1)
    for nome in ("P0", "P1", "P2", "P3"):
        linha += 1
        fundo, texto = COR_FAIXA[nome]
        c = ws.cell(row=linha, column=1, value=nome)
        c.font = Font(name=FONTE, bold=True, color=texto)
        c.fill = PatternFill("solid", fgColor=fundo)
        ws.cell(row=linha, column=2, value=DESCRICAO_FAIXA[nome]).font = normal
        ws.cell(row=linha, column=3, value=f"=COUNTIF('Por CVE'!${col_prio_cve}:${col_prio_cve},A{linha})")
        ws.cell(row=linha, column=4, value=f"=COUNTIF('Detalhe'!${col_prio_det}:${col_prio_det},A{linha})")
        for j in range(1, 5):
            ws.cell(row=linha, column=j).border = borda
            if j > 1:
                ws.cell(row=linha, column=j).font = normal
            ws.cell(row=linha, column=j).alignment = Alignment(wrap_text=True, vertical="center")
    linha += 1
    ws.cell(row=linha, column=1, value="Total").font = negrito
    ws.cell(row=linha, column=3, value=f"=SUM(C{inicio}:C{linha - 1})").font = negrito
    ws.cell(row=linha, column=4, value=f"=SUM(D{inicio}:D{linha - 1})").font = negrito
    for j in range(1, 5):
        ws.cell(row=linha, column=j).border = borda

    # --- variação vs. execução anterior ---
    if ctx.get("novos") is not None:
        linha += 2
        ws.cell(row=linha, column=1, value="CVEs novos").font = negrito
        ws.cell(row=linha, column=2, value=ctx["novos"]).font = normal
        linha += 1
        ws.cell(row=linha, column=1, value="CVEs resolvidos").font = negrito
        ws.cell(row=linha, column=2, value=ctx["resolvidos"]).font = normal
        for r in (linha - 1, linha):
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="left")

    # --- top 10 ---
    linha += 2
    ws.cell(row=linha, column=1, value="Top 10 CVEs").font = Font(
        name=FONTE, bold=True, size=12, color="1F3864"
    )
    linha += 1
    for j, titulo in enumerate(["CVE", "Motivo", "Faixa", "Agentes"], start=1):
        c = ws.cell(row=linha, column=j, value=titulo)
        c.font, c.fill, c.border = cab_fonte, cab_fill, borda
    for _, r in por_cve.head(10).iterrows():
        linha += 1
        valores = [r["cve"], r["motivo_prioridade"], r["prioridade"], int(r["agentes_afetados"])]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=j, value=v)
            c.font, c.border = normal, borda
            c.alignment = Alignment(wrap_text=True, vertical="top")
        fundo, texto = COR_FAIXA[r["prioridade"]]
        ws.cell(row=linha, column=3).fill = PatternFill("solid", fgColor=fundo)
        ws.cell(row=linha, column=3).font = Font(name=FONTE, bold=True, color=texto)

    # --- avisos ---
    linha += 2
    ws.cell(row=linha, column=1, value="Avisos de qualidade dos dados").font = Font(
        name=FONTE, bold=True, size=12, color="C00000"
    )
    if ctx["avisos"]:
        for aviso in ctx["avisos"]:
            linha += 1
            ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=4)
            c = ws.cell(row=linha, column=1, value="• " + aviso)
            c.font = normal
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[linha].height = 15 * (1 + len(aviso) // 110)
    else:
        linha += 1
        ws.cell(row=linha, column=1, value="Nenhum aviso.").font = normal

    # --- regras ---
    linha += 2
    ws.cell(row=linha, column=1, value="Regras aplicadas").font = Font(
        name=FONTE, bold=True, size=12, color="1F3864"
    )
    regras = [
        "P0: CVE está no CISA KEV.",
        f"P1: EPSS ≥ {REGRAS['p1_epss']} ou percentil ≥ {REGRAS['p1_percentil']}, "
        f"ou CVSS ≥ {REGRAS['p1_cvss_alto']} com EPSS ≥ {REGRAS['p1_cvss_alto_epss']}.",
        f"P2: CVSS ≥ {REGRAS['p2_cvss']} ou percentil ≥ {REGRAS['p2_percentil']}.",
        "P3: demais.",
        "Contexto do ativo: pode subir ou descer 1 nível. Só KEV é P0; KEV nunca cai abaixo de P1.",
        "Desempate dentro da faixa: 40% CVSS + 50% percentil EPSS + 10% uso em ransomware (KEV).",
        "'Sem EPSS' (CVE recente ou não encontrado) NÃO é tratado como baixo risco: a faixa usa só KEV/CVSS.",
        "Remediação: 'CVEs eliminados' assume atualização para uma versão corrigida; "
        "confirme se ela existe antes de abrir a mudança.",
    ]
    for texto in regras:
        linha += 1
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=4)
        c = ws.cell(row=linha, column=1, value=texto)
        c.font = normal
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[linha].height = 15 * (1 + len(texto) // 110)

    ws.sheet_view.showGridLines = False


def gerar_xlsx(caminho, ctx, por_cve, remediacao, por_ativo, detalhe, tendencia):
    abas = [
        ("Por CVE", por_cve),
        ("Remediação", remediacao),
        ("Por Ativo", por_ativo),
        ("Detalhe", detalhe),
        ("Tendência", tendencia),
    ]
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        # fonte padrão do workbook = Arial (evita estilizar célula a célula)
        writer.book._fonts = IndexedList([Font(name=FONTE, sz=10, family=2)])

        writer.book.create_sheet("Resumo")  # fica na primeira posição
        for nome, tabela in abas:
            _para_excel(tabela).to_excel(writer, sheet_name=nome, index=False)
        # remove a aba padrão criada pelo pandas se existir vazia
        if "Sheet" in writer.book.sheetnames:
            del writer.book["Sheet"]

        for nome, tabela in abas:
            _formatar_aba(writer.sheets[nome], tabela, largura_max=80 if nome == "Por CVE" else 60)
        _montar_resumo(writer.sheets["Resumo"], ctx, por_cve, detalhe)
        writer.book.active = 0


# =============================================================================
# RELATÓRIO HTML
# =============================================================================


def _html_escape(valor):
    import html
    if valor is None or (isinstance(valor, float) and np.isnan(valor)):
        return "—"
    texto = str(valor).strip()
    return html.escape(texto) if texto else "—"


def _tabela_html(df, colunas=None, limite=None):
    tabela = df.copy()
    if colunas is not None:
        tabela = tabela[[c for c in colunas if c in tabela.columns]]
    if limite is not None:
        tabela = tabela.head(limite)
    if tabela.empty:
        return '<p class="muted">Sem registros para esta seção.</p>'
    return tabela.to_html(index=False, border=0, classes="data-table", escape=True, na_rep="—")


def gerar_relatorio_html(caminho, ctx, por_cve, remediacao, por_ativo, detalhe, tendencia):
    """Gera relatório HTML local e autocontido; não envia dados a terceiros."""
    cont = por_cve["prioridade"].value_counts() if "prioridade" in por_cve else {}
    total_cves = int(por_cve["cve"].nunique()) if "cve" in por_cve else 0
    total_ocorrencias = int(len(detalhe))
    total_ativos = int(detalhe["agente"].nunique()) if "agente" in detalhe else 0
    total_pacotes = int(detalhe["pacote"].nunique()) if "pacote" in detalhe else 0
    kev_count = int(por_cve["in_kev"].fillna(False).astype(bool).sum()) if "in_kev" in por_cve else 0
    sem_epss_count = int(por_cve["sem_epss"].fillna(False).astype(bool).sum()) if "sem_epss" in por_cve else 0

    cards = "".join(f'<div class="metric"><span>{label}</span><strong>{value}</strong></div>' for label, value in [
        ("CVEs únicos", total_cves), ("Ocorrências", total_ocorrencias),
        ("Ativos afetados", total_ativos), ("Pacotes afetados", total_pacotes)])
    prioridades = "".join(
        f'<div class="priority {p.lower()}"><span>{p}</span><strong>{int(cont.get(p, 0))}</strong></div>'
        for p in ("P0", "P1", "P2", "P3"))
    avisos_html = ("<ul>" + "".join(f"<li>{_html_escape(a)}</li>" for a in ctx.get("avisos", [])) + "</ul>"
                   if ctx.get("avisos") else '<p class="ok">Nenhum aviso de qualidade registrado pelo processamento.</p>')

    top_cves = _tabela_html(por_cve, ["cve", "prioridade", "motivo_prioridade", "cvss", "epss", "epss_percentil", "in_kev", "agentes_afetados", "pacotes_afetados"], 50)
    rem_html = _tabela_html(remediacao, ["pacote", "versao_instalada", "prioridade_max", "hosts_afetados", "cves_eliminados", "P0", "P1", "P2", "P3", "cves_p0_p1", "condicao_correcao"], 100)
    ativo_html = _tabela_html(por_ativo, ["agente", "criticidade", "exposicao", "ambiente", "P0", "P1", "P2", "P3", "cves_total", "pacotes_vulneraveis"], 100)
    trend_html = _tabela_html(tendencia, ["data", "P0", "P1", "P2", "P3", "total_cves"], 180)

    detalhes_cve = []
    for _, r in por_cve.iterrows():
        cve = _html_escape(r.get("cve"))
        prioridade = _html_escape(r.get("prioridade"))
        descricao = _html_escape(r.get("descricao"))
        motivo = _html_escape(r.get("motivo_prioridade"))
        campos = [
            ("CVSS", r.get("cvss")), ("EPSS", r.get("epss")), ("Percentil EPSS", r.get("epss_percentil")),
            ("Presente no CISA KEV", "Sim" if bool(r.get("in_kev", False)) else "Não confirmado na consulta"),
            ("Uso conhecido em ransomware (KEV)", "Sim" if bool(r.get("kev_ransomware", False)) else "Não indicado"),
            ("Prazo CISA", r.get("kev_prazo")), ("Publicado no Wazuh", r.get("publicado")),
            ("Idade (dias)", r.get("idade_dias")), ("Agentes afetados", r.get("agentes_afetados")),
            ("Pacotes afetados", r.get("pacotes_afetados")),
            ("Novo no histórico", "Sim" if bool(r.get("novo", False)) else "Não / não disponível")]
        grid = "".join(f'<div><span>{_html_escape(k)}</span><strong>{_html_escape(v)}</strong></div>' for k, v in campos)
        detalhes_cve.append(f"""<article class="cve-card"><div class="cve-heading"><h3>{cve}</h3><span class="pill {prioridade.lower()}">{prioridade}</span></div>
        <p class="reason"><strong>Motivo da prioridade:</strong> {motivo}</p><p><strong>Descrição disponível no CSV do Wazuh:</strong> {descricao}</p>
        <div class="detail-grid">{grid}</div></article>""")
    detalhes_html = "\n".join(detalhes_cve) if detalhes_cve else '<p class="muted">Nenhum CVE para detalhar.</p>'

    css = """
    :root { color-scheme: light; --brand-orange:#f5a000; --brand-gray:#595959; --navy:#595959; --ink:#292929; --muted:#686868; --line:#dedede; --bg:#f4f4f4; --soft-orange:#fff3d9; }
    * { box-sizing:border-box; } body { margin:0; font-family:Arial,Helvetica,sans-serif; color:var(--ink); background:var(--bg); line-height:1.45; }
    .wrap { max-width:1240px; margin:0 auto; padding:28px; } header { background:var(--brand-gray); color:#fff; padding:30px; border-radius:10px; border-bottom:7px solid var(--brand-orange); }
    header h1 { margin:0 0 8px; font-size:27px; } header p { margin:4px 0; color:#f1f1f1; } header h1::first-letter { color:var(--brand-orange); }
    section { background:#fff; border:1px solid var(--line); border-radius:12px; padding:22px; margin-top:18px; box-shadow:0 2px 8px #162b4d0a; }
    h2 { margin:0 0 14px; color:var(--brand-gray); font-size:20px; border-left:4px solid var(--brand-orange); padding-left:10px; } h3 { margin:0; font-size:17px; }
    .metrics,.priorities { display:grid; grid-template-columns:repeat(4,minmax(0,1fr)); gap:12px; margin-top:18px; }
    .metric,.priority { padding:16px; border:1px solid var(--line); border-top:3px solid var(--brand-orange); border-radius:7px; background:#fff; display:flex; flex-direction:column; gap:5px; }
    .metric span,.detail-grid span { color:var(--muted); font-size:12px; } .metric strong { font-size:26px; }
    .priority { flex-direction:row; justify-content:space-between; align-items:center; font-weight:bold; } .priority strong { font-size:23px; }
    .p0 { background:#fde8e8; color:#a40000; } .p1 { background:#fff0df; color:#8a3b00; } .p2 { background:#fff8d9; color:#6d5700; } .p3 { background:#e6f4e8; color:#236334; }
    .meta { display:grid; grid-template-columns:200px 1fr; gap:8px 14px; } .meta b { color:var(--navy); }
    .data-table { border-collapse:collapse; width:100%; font-size:12px; margin-top:10px; } .data-table th { text-align:left; background:#595959; color:#fff; border-bottom:3px solid #f5a000; }
    .data-table th,.data-table td { padding:9px; border:1px solid var(--line); vertical-align:top; overflow-wrap:anywhere; } .data-table tr:nth-child(even) td { background:#fafbfd; }
    .table-scroll { overflow-x:auto; } .cve-card { border:1px solid var(--line); border-radius:10px; padding:18px; margin:12px 0; break-inside:avoid; }
    .cve-heading { display:flex; justify-content:space-between; align-items:center; gap:12px; } .pill { border-radius:20px; padding:5px 12px; font-size:12px; font-weight:bold; white-space:nowrap; }
    .reason { background:var(--soft-orange); border-left:3px solid var(--brand-orange); padding:10px; border-radius:4px; } .detail-grid { display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:10px; margin-top:14px; }
    .detail-grid div { background:#f8fafc; border:1px solid #e7ebf1; padding:10px; border-radius:7px; min-width:0; } .detail-grid span,.detail-grid strong { display:block; overflow-wrap:anywhere; }
    .muted { color:var(--muted); } .ok { color:#236334; } .notice { border-left:4px solid var(--brand-orange); padding:10px 14px; background:var(--soft-orange); }
    footer { color:var(--muted); font-size:12px; padding:20px 4px; } @media(max-width:700px) { .wrap{padding:12px}.metrics,.priorities{grid-template-columns:repeat(2,minmax(0,1fr))}.detail-grid{grid-template-columns:repeat(2,minmax(0,1fr))}.meta{grid-template-columns:1fr} }
    @media print { body{background:#fff}.wrap{max-width:none;padding:0} header,section{box-shadow:none} .cve-card{break-inside:avoid} .data-table{font-size:9px} }
    """
    html_doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
    <title>Relatório de Vulnerabilidades - {_html_escape(ctx.get('gerado_em'))}</title><style>{css}</style></head><body><main class="wrap">
    <header><h1>Relatório de Gestão de Vulnerabilidades</h1><p>Wazuh · EPSS (FIRST) · CISA KEV</p><p>Gerado em {_html_escape(ctx.get('gerado_em'))}</p><p>Documento técnico para apoio à priorização e remediação.</p></header>
    <div class="metrics">{cards}</div><div class="priorities">{prioridades}</div>
    <section><h2>1. Resumo executivo</h2><p>Foram identificados <strong>{total_cves}</strong> CVEs únicos em <strong>{total_ativos}</strong> ativos e <strong>{total_pacotes}</strong> pacotes distintos. A classificação reflete as regras configuradas e os dados disponíveis nesta execução.</p>
    <p><strong>CVEs listados no KEV:</strong> {kev_count}. <strong>CVEs sem score EPSS:</strong> {sem_epss_count}.</p><div class="notice">A prioridade é um indicador de triagem. Confirme produto, versão afetada, exposição e existência de correção antes de aprovar a remediação. Ausência no KEV não prova ausência de exploração.</div></section>
    <section><h2>2. Fontes e contexto do processamento</h2><div class="meta"><b>Arquivo de entrada</b><span>{_html_escape(ctx.get('entrada'))}</span><b>Dados EPSS</b><span>{_html_escape(ctx.get('epss_info'))}</span><b>Catálogo CISA KEV</b><span>{_html_escape(ctx.get('kev_info'))}</span><b>Contexto de ativos</b><span>{_html_escape(ctx.get('ativos_info'))}</span><b>Execução anterior</b><span>{_html_escape(ctx.get('anterior_info'))}</span><b>CVEs novos</b><span>{_html_escape(ctx.get('novos'))}</span><b>CVEs resolvidos</b><span>{_html_escape(ctx.get('resolvidos'))}</span></div></section>
    <section><h2>3. CVEs priorizados</h2><p class="muted">Exibe até 50 CVEs por prioridade. Consulte o XLSX/CSV para a listagem completa.</p><div class="table-scroll">{top_cves}</div></section>
    <section><h2>4. Análise individual por CVE</h2><p class="muted">A descrição é baseada no campo do CSV do Wazuh; os demais dados vêm do processamento atual. Campos indisponíveis aparecem como traço.</p>{detalhes_html}</section>
    <section><h2>5. Plano de remediação por pacote</h2><div class="table-scroll">{rem_html}</div></section><section><h2>6. Visão por ativo</h2><div class="table-scroll">{ativo_html}</div></section>
    <section><h2>7. Tendência histórica</h2><div class="table-scroll">{trend_html}</div></section><section><h2>8. Qualidade dos dados e avisos</h2>{avisos_html}<p>Este relatório não consulta individualmente a NVD. Descrição e aplicabilidade dependem do CSV de origem; EPSS e KEV dependem do estado das fontes/caches indicado acima.</p></section>
    <section><h2>9. Critérios de classificação</h2><ul><li><strong>P0:</strong> CVE consta no CISA KEV.</li><li><strong>P1:</strong> EPSS ≥ {REGRAS['p1_epss']:.2f}, percentil EPSS ≥ {REGRAS['p1_percentil']:.0%}, ou CVSS ≥ {REGRAS['p1_cvss_alto']:.1f} com EPSS ≥ {REGRAS['p1_cvss_alto_epss']:.2f}.</li><li><strong>P2:</strong> CVSS ≥ {REGRAS['p2_cvss']:.1f} ou percentil EPSS ≥ {REGRAS['p2_percentil']:.0%}.</li><li><strong>P3:</strong> demais casos, conforme regras atuais.</li><li>O contexto pode ajustar a prioridade em até um nível. CVE KEV não cai abaixo de P1.</li></ul></section>
    <footer>Relatório gerado localmente. Trate como informação interna: restrinja o acesso, revise antes de compartilhar e aplique a política de retenção da organização.</footer></main></body></html>"""
    caminho.write_text(html_doc, encoding="utf-8")


# =============================================================================
# PRINCIPAL
# =============================================================================


def parse_args():
    p = argparse.ArgumentParser(description="Priorização de vulnerabilidades Wazuh + EPSS + KEV")
    p.add_argument("--entrada", type=Path, default=Path("input/vulns.csv"))
    p.add_argument("--ativos", type=Path, default=Path("input/ativos.csv"))
    p.add_argument("--saida-dir", type=Path, default=Path("output"))
    p.add_argument("--cache-dir", type=Path, default=Path("cache"))
    p.add_argument("--sem-html", action="store_true", help="não gerar o relatório HTML")
    return p.parse_args()


def main():
    args = parse_args()
    hoje = date.today().isoformat()
    hoje_dt = pd.Timestamp(hoje)
    avisos = []

    # 1) entrada
    df, originais = carregar_wazuh(args.entrada, avisos)
    df["_idx"] = df.index
    cves_validos = sorted(df.loc[df["cve_valido"], "cve"].unique().tolist())
    log(f"CVEs válidos únicos: {len(cves_validos)}")
    if not cves_validos:
        raise ValueError("Nenhum CVE válido foi encontrado no CSV.")

    # 2) fontes externas
    epss, epss_data, epss_origem = obter_epss(cves_validos, args.cache_dir, hoje, avisos)
    kev, kev_ok = obter_kev(args.cache_dir, hoje, avisos)
    ativos = carregar_ativos(args.ativos, avisos)

    sem_epss = [c for c in cves_validos if c not in epss.index]
    if sem_epss:
        avisos.append(
            f"{len(sem_epss)} de {len(cves_validos)} CVEs sem EPSS (recentes ou ainda não "
            "pontuados). Priorizados por KEV/CVSS, sem assumir baixo risco."
        )
    if ativos is not None:
        sem_ctx = sorted(set(df["agente"].map(normalizar)) - set(ativos["_chave"]))
        if sem_ctx:
            avisos.append(
                f"{len(sem_ctx)} agentes sem contexto em ativos.csv (sem ajuste): "
                + ", ".join(sem_ctx[:8])
                + ("..." if len(sem_ctx) > 8 else "")
            )

    # 3) histórico
    hist_dir = args.saida_dir / "historico"
    anterior, data_anterior = carregar_snapshot_anterior(hist_dir, hoje)

    # 4) priorização e visões
    log("Calculando prioridades ...")
    df = priorizar(df, epss, kev, kev_ok, ativos, anterior, hoje_dt)
    por_cve = visao_por_cve(df, anterior)
    remediacao = visao_remediacao(df)
    por_ativo = visao_por_ativo(df)
    detalhe = visao_detalhe(df, originais)

    # 5) snapshot e tendência
    hist_dir.mkdir(parents=True, exist_ok=True)
    por_cve[["cve", "prioridade", "epss", "epss_percentil", "in_kev", "cvss"]].to_csv(
        hist_dir / f"cves_{hoje}.csv", index=False
    )
    tendencia = montar_tendencia(hist_dir)

    ctx = {
        "gerado_em": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "entrada": args.entrada,
        "epss_info": f"{len(epss)} CVEs na base; score de {epss_data}; origem: {epss_origem}",
        "kev_info": (
            f"{len(kev)} CVEs no catálogo; {int(por_cve['in_kev'].sum())} presentes neste CSV"
            if kev_ok
            else "INDISPONÍVEL nesta execução"
        ),
        "ativos_info": (
            f"{args.ativos} ({len(ativos)} agentes mapeados)" if ativos is not None else "não utilizado"
        ),
        "anterior_info": data_anterior or "nenhuma (primeira execução)",
        "avisos": avisos,
        "novos": None,
        "resolvidos": None,
    }
    if anterior is not None:
        atuais = set(por_cve["cve"])
        antigos = set(anterior.index)
        ctx["novos"] = len(atuais - antigos)
        ctx["resolvidos"] = len(antigos - atuais)

    # 6) saídas
    args.saida_dir.mkdir(parents=True, exist_ok=True)
    base = args.saida_dir / f"priorizacao_{hoje}"
    gerar_xlsx(base.with_suffix(".xlsx"), ctx, por_cve, remediacao, por_ativo, detalhe, tendencia)
    detalhe.to_csv(base.with_suffix(".csv"), index=False, encoding="utf-8-sig")
    por_cve.to_json(base.with_suffix(".json"), orient="records", force_ascii=False, indent=2)
    relatorio_html = base.with_name(f"relatorio_vulnerabilidades_{hoje}").with_suffix(".html")
    if not args.sem_html:
        gerar_relatorio_html(relatorio_html, ctx, por_cve, remediacao, por_ativo, detalhe, tendencia)

    # 7) resumo no terminal
    cont = por_cve["prioridade"].value_counts()
    print("\nProcessamento concluído!")
    print(f"Registros (agente x pacote x CVE): {len(df)}")
    print(f"CVEs únicos: {len(por_cve)}  |  "
          + "  ".join(f"{p}: {int(cont.get(p, 0))}" for p in NIVEIS.values()))
    print(f"CVEs sem EPSS: {len(sem_epss)}")
    for aviso in avisos:
        print(f"  [aviso] {aviso}")
    print(f"Arquivos: {base}.xlsx / .csv / .json")
    if not args.sem_html:
        print(f"Relatório HTML: {relatorio_html}")


if __name__ == "__main__":
    main()
